"""PDF and DOCX text extraction with page-level mapping.

Each parsed page returns a dict:
    {page: int (0-based), text: str}

For PDFs we extract embedded text via PyMuPDF; pages with no text are flagged
as "scanned" and sent to OCR (if enabled). For DOCX we treat the whole
document as a single "page 0" since Word files don't have fixed pages.
"""
from __future__ import annotations

import logging
from pathlib import Path

from app.ingest.ocr import ocr_page

logger = logging.getLogger("app")


class ParseError(Exception):
    pass


def _pdf_page_to_image(doc, page_index: int) -> bytes | None:
    """Render a PyMuPDF page to PNG bytes (for OCR). None on failure."""
    try:
        page = doc.load_page(page_index)
        pix = page.get_pixmap(dpi=200)
        return pix.tobytes("png")
    except Exception:  # noqa: BLE001
        logger.exception("Failed to rasterize page %s for OCR", page_index)
        return None


def parse_pdf(path: str | Path, *, use_ocr: bool = True) -> list[dict]:
    """Extract text per PDF page.

    Returns list of {"page": int, "text": str}. Pages with no embedded text
    are OCR'd when use_ocr is True and an OCR backend is available.
    """
    import fitz  # PyMuPDF  # noqa: PLC0415

    path = Path(path)
    result: list[dict] = []
    try:
        doc = fitz.open(str(path))
        total = doc.page_count
        for idx in range(total):
            page = doc.load_page(idx)
            text = page.get_text("text").strip()
            if not text and use_ocr:
                img = _pdf_page_to_image(doc, idx)
                if img:
                    ocr_text = ocr_page(img)
                    if ocr_text:
                        text = ocr_text
                        logger.info("OCR applied to scanned page %d of %s", idx + 1, path.name)
            result.append({"page": idx, "text": text})
        doc.close()
        return result
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Failed to parse PDF {path.name}: {exc}") from exc


def parse_docx(path: str | Path) -> list[dict]:
    """Extract text from a Word document (single logical 'page').

    Includes paragraph text and table cells, separated clearly.
    """
    from docx import Document  # noqa: PLC0415

    path = Path(path)
    try:
        doc = Document(str(path))
        parts: list[str] = []

        # Body paragraphs
        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                parts.append(t)

        # Tables (preserve structure)
        for ti, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                parts.append("[TABLE %d]\n%s" % (ti, "\n".join(rows)))

        text = "\n".join(parts).strip()
        if not text:
            raise ParseError(f"No text extracted from DOCX {path.name}")
        return [{"page": 0, "text": text}]
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Failed to parse DOCX {path.name}: {exc}") from exc


def parse_file(path: str | Path, *, use_ocr: bool = True) -> tuple[list[dict], str]:
    """Parse a file into per-page text. Returns (pages, file_type).

    file_type is "pdf", "docx" or "txt".

    `use_ocr` controls whether scanned/image PDF pages are run through OCR.
    - ``True`` (default): pages with no embedded text are OCR'd (if a backend
      is available).
    - ``False`` (direct/embedding-only indexing): PDFs are extracted for text
      only; image-only pages are left empty (faster, no OCR cost).

    DOCX / TXT / MD already contain their text directly, so `use_ocr` has no
    effect on them (it applies to scanned PDF pages).
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path, use_ocr=use_ocr), "pdf"
    if suffix in (".docx", ".doc"):
        return parse_docx(path), "docx"
    if suffix in (".txt", ".md", ".markdown", ".text"):
        return parse_plain_text(path), "txt"
    raise ParseError(f"Unsupported file type: {suffix or '(none)'}. Use PDF, DOCX, TXT or Markdown.")


def parse_plain_text(path: str | Path) -> list[dict]:
    """Read a plain-text / Markdown file, splitting into pages by blank lines.

    Keeps the text intact (no markdown stripping) so your notes stay readable.
    """
    path = Path(path)
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Could not read text file {path.name}: {exc}") from exc
    raw = raw.strip()
    if not raw:
        raise ParseError(f"No text extracted from {path.name}")
    # Split into logical "pages" on two+ blank lines for better chunk mapping.
    import re

    parts = [p.strip() for p in re.split(r"\n\s*\n\s*\n+", raw) if p.strip()]
    return [{"page": i, "text": part} for i, part in enumerate(parts)]
